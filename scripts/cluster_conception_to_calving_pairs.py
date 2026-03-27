from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference


BIN_EDGES = [-np.inf, 260, 270, 280, 290, 300, 310, 320, 330, 340, np.inf]
BIN_LABELS = [
    "<260",
    "260-269",
    "270-279",
    "280-289",
    "290-299",
    "300-309",
    "310-319",
    "320-329",
    "330-339",
    "340+",
]


def _fill_text(s: pd.Series, default: str = "unknown") -> pd.Series:
    return s.astype("string").fillna(default).replace("", default)


def _collapse_rare(s: pd.Series, *, min_rows: int, other_label: str = "OTHER") -> pd.Series:
    filled = _fill_text(s)
    counts = filled.value_counts(dropna=False)
    keep = set(counts[counts >= min_rows].index.tolist())
    return filled.where(filled.isin(keep), other_label)


def _zscore(s: pd.Series) -> np.ndarray:
    x = pd.to_numeric(s, errors="coerce").astype(float)
    mean = np.nanmean(x)
    std = np.nanstd(x)
    if not np.isfinite(std) or std <= 1e-12:
        return np.zeros(len(x), dtype=float)
    x = np.where(np.isfinite(x), x, mean)
    return (x - mean) / std


def _prepare_frame(path: str | Path) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name="all_pairs")
    df = df.copy()
    df["days_from_conception_to_calving"] = pd.to_numeric(df["days_from_conception_to_calving"], errors="coerce")
    df = df[df["days_from_conception_to_calving"].notna()].copy()
    df["days_from_conception_to_calving"] = df["days_from_conception_to_calving"].astype(int)

    for col in ("conception_date", "calving_date"):
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce")

    df["gestation_bin_10d"] = pd.cut(
        df["days_from_conception_to_calving"],
        bins=BIN_EDGES,
        labels=BIN_LABELS,
        right=False,
        include_lowest=True,
    ).astype("string")
    df["gestation_tail"] = np.where(
        df["days_from_conception_to_calving"] < 270,
        "short",
        np.where(df["days_from_conception_to_calving"] > 320, "long", "core"),
    )

    df["conception_subdivision_group"] = _collapse_rare(df["conception_subdivision"], min_rows=100)
    df["conception_season_group"] = _fill_text(df["conception_season"])
    df["lactation_group_group"] = _fill_text(df["lactation_group"])
    df["heifer_group"] = np.where(df["is_heifer_pregnancy"].fillna(False), "heifer", "cow")
    df["semen_type_group"] = np.where(
        _fill_text(df.get("semen_type", pd.Series(index=df.index, dtype="string"))).str.lower().eq("sex"),
        "sex",
        "traditional",
    )
    df["bull_breed_group"] = _collapse_rare(df.get("bull_breed", pd.Series(index=df.index, dtype="string")), min_rows=200)
    df["bull_plem_group"] = _collapse_rare(
        df.get("bull_plem", pd.Series(index=df.index)).astype("string"),
        min_rows=100,
    )
    calf_count = pd.to_numeric(df.get("calf_count_in_calving", pd.Series(index=df.index)), errors="coerce")
    df["calf_count_group"] = np.where(
        calf_count.isna(),
        "unknown",
        np.where(calf_count <= 1, "1", np.where(calf_count == 2, "2", "3+")),
    )
    is_multi = df.get("is_multiple_birth", pd.Series(index=df.index)).fillna(0).astype(float) > 0
    df["multiple_birth_group"] = np.where(is_multi, "multiple", "single")
    df["dim_or_age_at_conception"] = pd.to_numeric(df["dim_or_age_at_conception"], errors="coerce")
    return df


def _build_feature_matrix(df: pd.DataFrame) -> tuple[np.ndarray, list[str]]:
    blocks: list[np.ndarray] = []
    names: list[str] = []

    categorical_features = [
        "conception_subdivision_group",
        "conception_season_group",
        "lactation_group_group",
        "heifer_group",
        "semen_type_group",
        "bull_breed_group",
        "bull_plem_group",
        "calf_count_group",
        "multiple_birth_group",
    ]

    for feature in categorical_features:
        dummies = pd.get_dummies(_fill_text(df[feature]), prefix=feature)
        if dummies.shape[1] == 0:
            continue
        weight = 1.0 / np.sqrt(float(dummies.shape[1]))
        block = dummies.to_numpy(dtype=float) * weight
        blocks.append(block)
        names.extend(dummies.columns.tolist())

    numeric_specs = [
        ("dim_or_age_at_conception", 1.0),
        ("days_from_conception_to_calving", 1.5),
    ]
    for col, weight in numeric_specs:
        arr = _zscore(df[col]).reshape(-1, 1) * weight
        blocks.append(arr)
        names.append(col)

    return np.hstack(blocks), names


def _kmeans_pp_init(x: np.ndarray, n_clusters: int, rng: np.random.Generator) -> np.ndarray:
    n_samples = x.shape[0]
    centers = np.empty((n_clusters, x.shape[1]), dtype=float)
    first_idx = int(rng.integers(0, n_samples))
    centers[0] = x[first_idx]
    closest_sq = np.sum((x - centers[0]) ** 2, axis=1)
    for i in range(1, n_clusters):
        probs = closest_sq / max(closest_sq.sum(), 1e-12)
        idx = int(rng.choice(n_samples, p=probs))
        centers[i] = x[idx]
        dist_sq = np.sum((x - centers[i]) ** 2, axis=1)
        closest_sq = np.minimum(closest_sq, dist_sq)
    return centers


def _run_kmeans(
    x: np.ndarray,
    *,
    n_clusters: int,
    random_state: int = 42,
    n_init: int = 8,
    max_iter: int = 100,
) -> tuple[np.ndarray, np.ndarray, float]:
    rng = np.random.default_rng(random_state)
    best_labels = None
    best_centers = None
    best_inertia = None

    for _ in range(n_init):
        centers = _kmeans_pp_init(x, n_clusters, rng)
        labels = np.zeros(x.shape[0], dtype=int)
        for _ in range(max_iter):
            dist = np.sum((x[:, None, :] - centers[None, :, :]) ** 2, axis=2)
            new_labels = np.argmin(dist, axis=1)
            if np.array_equal(new_labels, labels):
                break
            labels = new_labels
            new_centers = centers.copy()
            for k in range(n_clusters):
                mask = labels == k
                if mask.any():
                    new_centers[k] = x[mask].mean(axis=0)
                else:
                    new_centers[k] = x[int(rng.integers(0, x.shape[0]))]
            centers = new_centers
        inertia = float(np.sum((x - centers[labels]) ** 2))
        if best_inertia is None or inertia < best_inertia:
            best_inertia = inertia
            best_labels = labels.copy()
            best_centers = centers.copy()

    assert best_labels is not None and best_centers is not None and best_inertia is not None
    return best_labels, best_centers, best_inertia


def _cluster_metric_summary(df: pd.DataFrame) -> pd.DataFrame:
    grouped = (
        df.groupby("cluster_id", dropna=False)["days_from_conception_to_calving"]
        .agg(["count", "mean", "median", "std", "min", "max"])
        .reset_index()
        .rename(
            columns={
                "count": "rows",
                "mean": "mean_days",
                "median": "median_days",
                "std": "std_days",
                "min": "min_days",
                "max": "max_days",
            }
        )
    )
    grouped["share_pct"] = grouped["rows"] / max(len(df), 1) * 100.0
    q = (
        df.groupby("cluster_id", dropna=False)["days_from_conception_to_calving"]
        .agg(
            p10_days=lambda s: s.quantile(0.10),
            p90_days=lambda s: s.quantile(0.90),
        )
        .reset_index()
    )
    grouped = grouped.merge(q, on="cluster_id", how="left")
    grouped["mean_dim_or_age"] = (
        df.groupby("cluster_id", dropna=False)["dim_or_age_at_conception"].mean().reset_index(drop=True)
    )
    grouped["heifer_share_pct"] = (
        df.groupby("cluster_id", dropna=False)["is_heifer_pregnancy"].mean().reset_index(drop=True) * 100.0
    )
    grouped["multiple_birth_share_pct"] = (
        df.groupby("cluster_id", dropna=False)["multiple_birth_group"].apply(lambda s: s.eq("multiple").mean()).reset_index(drop=True)
        * 100.0
    )
    grouped["sexed_semen_share_pct"] = (
        df.groupby("cluster_id", dropna=False)["semen_type_group"].apply(lambda s: s.eq("sex").mean()).reset_index(drop=True)
        * 100.0
    )
    return grouped.sort_values("mean_days").reset_index(drop=True)


def _dominant_by_cluster(df: pd.DataFrame, feature: str) -> pd.DataFrame:
    work = df[["cluster_id", feature]].copy()
    work[feature] = _fill_text(work[feature])
    counts = (
        work.groupby(["cluster_id", feature], dropna=False)
        .size()
        .rename("rows")
        .reset_index()
        .rename(columns={feature: "category"})
    )
    totals = counts.groupby("cluster_id", dropna=False)["rows"].sum().rename("cluster_rows").reset_index()
    counts = counts.merge(totals, on="cluster_id", how="left")
    counts["share_within_cluster_pct"] = counts["rows"] / counts["cluster_rows"] * 100.0
    counts = counts.sort_values(
        ["cluster_id", "share_within_cluster_pct", "rows", "category"],
        ascending=[True, False, False, True],
        kind="mergesort",
    )
    top = counts.groupby("cluster_id", dropna=False).head(1).copy()
    top = top.rename(
        columns={
            "category": f"dominant_{feature}",
            "share_within_cluster_pct": f"dominant_{feature}_share_pct",
        }
    )
    return top[["cluster_id", f"dominant_{feature}", f"dominant_{feature}_share_pct"]]


def _cluster_feature_mix(df: pd.DataFrame, feature: str) -> pd.DataFrame:
    work = df[["cluster_id", feature]].copy()
    work[feature] = _fill_text(work[feature])
    counts = (
        work.groupby(["cluster_id", feature], dropna=False)
        .size()
        .rename("rows")
        .reset_index()
        .rename(columns={feature: "category"})
    )
    totals = counts.groupby("cluster_id", dropna=False)["rows"].sum().rename("cluster_rows").reset_index()
    counts = counts.merge(totals, on="cluster_id", how="left")
    counts["share_within_cluster_pct"] = counts["rows"] / counts["cluster_rows"] * 100.0
    counts.insert(1, "feature", feature)
    return counts.sort_values(["cluster_id", "feature", "rows"], ascending=[True, True, False], kind="mergesort")


def _cluster_profiles(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    summary = _cluster_metric_summary(df)
    dominant_frames = [
        _dominant_by_cluster(df, "conception_subdivision_group"),
        _dominant_by_cluster(df, "conception_season_group"),
        _dominant_by_cluster(df, "lactation_group_group"),
        _dominant_by_cluster(df, "heifer_group"),
        _dominant_by_cluster(df, "semen_type_group"),
        _dominant_by_cluster(df, "bull_breed_group"),
        _dominant_by_cluster(df, "bull_plem_group"),
        _dominant_by_cluster(df, "calf_count_group"),
        _dominant_by_cluster(df, "multiple_birth_group"),
    ]
    for frame in dominant_frames:
        summary = summary.merge(frame, on="cluster_id", how="left")

    mix = pd.concat(
        [
            _cluster_feature_mix(df, "conception_subdivision_group"),
            _cluster_feature_mix(df, "conception_season_group"),
            _cluster_feature_mix(df, "lactation_group_group"),
            _cluster_feature_mix(df, "heifer_group"),
            _cluster_feature_mix(df, "semen_type_group"),
            _cluster_feature_mix(df, "bull_breed_group"),
            _cluster_feature_mix(df, "bull_plem_group"),
            _cluster_feature_mix(df, "calf_count_group"),
            _cluster_feature_mix(df, "multiple_birth_group"),
        ],
        ignore_index=True,
    )
    return summary, mix


def _dominant_by_bin(df: pd.DataFrame, feature: str) -> pd.DataFrame:
    work = df[["gestation_bin_10d", feature]].copy()
    work[feature] = _fill_text(work[feature])
    counts = (
        work.groupby(["gestation_bin_10d", feature], dropna=False)
        .size()
        .rename("rows")
        .reset_index()
        .rename(columns={feature: "category", "gestation_bin_10d": "bin"})
    )
    totals = counts.groupby("bin", dropna=False)["rows"].sum().rename("bin_rows").reset_index()
    counts = counts.merge(totals, on="bin", how="left")
    counts["share_within_bin_pct"] = counts["rows"] / counts["bin_rows"] * 100.0
    counts = counts.sort_values(
        ["bin", "share_within_bin_pct", "rows", "category"],
        ascending=[True, False, False, True],
        kind="mergesort",
    )
    top = counts.groupby("bin", dropna=False).head(1).copy()
    top = top.rename(
        columns={
            "category": f"dominant_{feature}",
            "share_within_bin_pct": f"dominant_{feature}_share_pct",
        }
    )
    return top[["bin", f"dominant_{feature}", f"dominant_{feature}_share_pct"]]


def _bin_profiles(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    summary = (
        df.groupby("gestation_bin_10d", dropna=False)["days_from_conception_to_calving"]
        .agg(["count", "mean", "median", "min", "max"])
        .reset_index()
        .rename(
            columns={
                "gestation_bin_10d": "bin",
                "count": "rows",
                "mean": "mean_days",
                "median": "median_days",
                "min": "min_days",
                "max": "max_days",
            }
        )
    )
    summary["share_pct"] = summary["rows"] / max(len(df), 1) * 100.0
    summary["mean_dim_or_age"] = (
        df.groupby("gestation_bin_10d", dropna=False)["dim_or_age_at_conception"].mean().reset_index(drop=True)
    )
    summary["heifer_share_pct"] = (
        df.groupby("gestation_bin_10d", dropna=False)["is_heifer_pregnancy"].mean().reset_index(drop=True) * 100.0
    )
    summary["multiple_birth_share_pct"] = (
        df.groupby("gestation_bin_10d", dropna=False)["multiple_birth_group"].apply(lambda s: s.eq("multiple").mean()).reset_index(drop=True)
        * 100.0
    )
    summary["sexed_semen_share_pct"] = (
        df.groupby("gestation_bin_10d", dropna=False)["semen_type_group"].apply(lambda s: s.eq("sex").mean()).reset_index(drop=True)
        * 100.0
    )

    dominant_frames = [
        _dominant_by_bin(df, "conception_subdivision_group"),
        _dominant_by_bin(df, "conception_season_group"),
        _dominant_by_bin(df, "lactation_group_group"),
        _dominant_by_bin(df, "heifer_group"),
        _dominant_by_bin(df, "semen_type_group"),
        _dominant_by_bin(df, "bull_breed_group"),
        _dominant_by_bin(df, "bull_plem_group"),
        _dominant_by_bin(df, "calf_count_group"),
        _dominant_by_bin(df, "multiple_birth_group"),
    ]
    for frame in dominant_frames:
        summary = summary.merge(frame, on="bin", how="left")

    mix_frames = []
    for feature in [
        "conception_subdivision_group",
        "conception_season_group",
        "lactation_group_group",
        "heifer_group",
        "semen_type_group",
        "bull_breed_group",
        "bull_plem_group",
        "calf_count_group",
        "multiple_birth_group",
    ]:
        work = df[["gestation_bin_10d", feature]].copy()
        work[feature] = _fill_text(work[feature])
        counts = (
            work.groupby(["gestation_bin_10d", feature], dropna=False)
            .size()
            .rename("rows")
            .reset_index()
            .rename(columns={"gestation_bin_10d": "bin", feature: "category"})
        )
        totals = counts.groupby("bin", dropna=False)["rows"].sum().rename("bin_rows").reset_index()
        counts = counts.merge(totals, on="bin", how="left")
        counts["share_within_bin_pct"] = counts["rows"] / counts["bin_rows"] * 100.0
        counts.insert(1, "feature", feature)
        mix_frames.append(counts)
    mix = pd.concat(mix_frames, ignore_index=True)
    return summary, mix


def _write_excel(
    *,
    out_path: Path,
    overall_summary: pd.DataFrame,
    cluster_summary: pd.DataFrame,
    cluster_feature_mix: pd.DataFrame,
    bin_profile_summary: pd.DataFrame,
    bin_feature_mix: pd.DataFrame,
    pairs_with_clusters: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        overall_summary.to_excel(writer, sheet_name="summary", index=False)
        cluster_summary.to_excel(writer, sheet_name="cluster_summary", index=False)
        cluster_feature_mix.to_excel(writer, sheet_name="cluster_feature_mix", index=False)
        bin_profile_summary.to_excel(writer, sheet_name="bin_profiles", index=False)
        bin_feature_mix.to_excel(writer, sheet_name="bin_feature_mix", index=False)
        pairs_with_clusters.to_excel(writer, sheet_name="pairs_with_cluster", index=False)


def _add_excel_charts(workbook_path: Path) -> None:
    wb = load_workbook(workbook_path)
    if "charts" in wb.sheetnames:
        del wb["charts"]
    ws = wb.create_sheet("charts")

    cluster_ws = wb["cluster_summary"]

    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "Размер кластеров"
    chart1.y_axis.title = "Количество строк"
    chart1.x_axis.title = "Кластер"
    data = Reference(cluster_ws, min_col=2, min_row=1, max_row=cluster_ws.max_row)
    cats = Reference(cluster_ws, min_col=1, min_row=2, max_row=cluster_ws.max_row)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    ws.add_chart(chart1, "A1")

    chart2 = BarChart()
    chart2.type = "col"
    chart2.style = 11
    chart2.title = "Средний срок по кластерам"
    chart2.y_axis.title = "Дней"
    chart2.x_axis.title = "Кластер"
    data = Reference(cluster_ws, min_col=3, min_row=1, max_row=cluster_ws.max_row)
    cats = Reference(cluster_ws, min_col=1, min_row=2, max_row=cluster_ws.max_row)
    chart2.add_data(data, titles_from_data=True)
    chart2.set_categories(cats)
    ws.add_chart(chart2, "J1")

    wb.save(workbook_path)


def analyze_file(
    input_path: str | Path,
    *,
    output_path: str | Path,
    n_clusters: int,
    random_state: int,
) -> Path:
    input_path = Path(input_path)
    output_path = Path(output_path)

    df = _prepare_frame(input_path)
    x, _feature_names = _build_feature_matrix(df)
    labels, _centers, inertia = _run_kmeans(
        x,
        n_clusters=n_clusters,
        random_state=random_state,
    )

    df = df.copy()
    df["cluster_id"] = labels.astype(int) + 1

    overall_summary = pd.DataFrame(
        [
            {
                "rows": int(len(df)),
                "clusters": int(n_clusters),
                "inertia": float(inertia),
                "mean_days": float(df["days_from_conception_to_calving"].mean()),
                "median_days": float(df["days_from_conception_to_calving"].median()),
                "short_share_pct": float((df["days_from_conception_to_calving"] < 270).mean() * 100.0),
                "core_share_pct": float(df["days_from_conception_to_calving"].between(270, 320, inclusive="both").mean() * 100.0),
                "long_share_pct": float((df["days_from_conception_to_calving"] > 320).mean() * 100.0),
            }
        ]
    )

    cluster_summary, cluster_feature_mix = _cluster_profiles(df)
    bin_profile_summary, bin_feature_mix = _bin_profiles(df)

    export_cols = [
        "cluster_id",
        "reg",
        "days_from_conception_to_calving",
        "gestation_bin_10d",
        "gestation_tail",
        "conception_date",
        "calving_date",
        "conception_subdivision",
        "conception_season",
        "lactation_at_conception",
        "lactation_group",
        "lactation_at_calving",
        "dim_or_age_at_conception",
        "is_heifer_pregnancy",
        "semen_type",
        "bull_breed",
        "bull_plem",
        "calf_count_in_calving",
        "male_calves_in_calving",
        "female_calves_in_calving",
        "is_multiple_birth",
    ]
    pairs_with_clusters = df[[c for c in export_cols if c in df.columns]].copy()

    _write_excel(
        out_path=output_path,
        overall_summary=overall_summary,
        cluster_summary=cluster_summary,
        cluster_feature_mix=cluster_feature_mix,
        bin_profile_summary=bin_profile_summary,
        bin_feature_mix=bin_feature_mix,
        pairs_with_clusters=pairs_with_clusters,
    )
    _add_excel_charts(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default="exports/conception_to_calving_pairs_enalb.xlsx",
        help="Path to conception-to-calving pairs Excel",
    )
    parser.add_argument(
        "--output",
        default="exports/conception_to_calving_pairs_enalb_clusters.xlsx",
        help="Path to output cluster Excel",
    )
    parser.add_argument(
        "--clusters",
        type=int,
        default=8,
        help="Number of clusters",
    )
    parser.add_argument(
        "--random-state",
        type=int,
        default=42,
        help="Random seed for k-means",
    )
    args = parser.parse_args()

    out = analyze_file(
        args.input,
        output_path=args.output,
        n_clusters=args.clusters,
        random_state=args.random_state,
    )
    print(f"saved_excel: {out}")


if __name__ == "__main__":
    main()
