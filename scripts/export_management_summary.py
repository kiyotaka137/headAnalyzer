from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage


MAIN_GROUPS = [
    "Дойные коровы",
    "Сухостойные коровы",
    "Нетели",
    "Тёлки 0–3 мес",
    "Тёлки 3–8 мес",
    "Тёлки ≥9 мес",
    "Бычки 0–2 мес",
    "Ожидаемый отёл, всего",
]

MONTH_RU = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}


def _fill_text(s: pd.Series, default: str = "не указано") -> pd.Series:
    return s.astype("string").fillna(default).replace("", default)


def _normalize_name_series(s: pd.Series) -> pd.Series:
    return (
        s.astype("string")
        .fillna("")
        .str.replace("\u00a0", " ", regex=False)
        .str.strip()
        .str.upper()
    )


def _extract_allowed_subdivisions(path: str | Path) -> list[str]:
    path = Path(path)
    xl = pd.ExcelFile(path)

    if "summary_long" in xl.sheet_names:
        df = pd.read_excel(path, sheet_name="summary_long", usecols=["Подразделение"])
        vals = (
            df["Подразделение"]
            .dropna()
            .astype("string")
            .str.strip()
        )
        return sorted(v for v in vals.unique().tolist() if v)

    if "Основные группы" in xl.sheet_names:
        df = pd.read_excel(path, sheet_name="Основные группы", usecols=["Подразделение"])
        vals = (
            df["Подразделение"]
            .dropna()
            .astype("string")
            .str.strip()
        )
        return sorted(v for v in vals.unique().tolist() if v)

    if "Сводка по прогнозу" in xl.sheet_names:
        df = pd.read_excel(path, sheet_name="Сводка по прогнозу", header=None)
        first_col = df.iloc[:, 0].dropna().astype("string").str.strip()
        vals = [
            str(v)
            for v in first_col.tolist()
            if v
            and v != "Подразделение"
            and ":" not in str(v)
            and not str(v).startswith("Все группы")
            and not str(v).startswith("Основные группы")
        ]
        return sorted(set(vals))

    return []


def _top_category(s: pd.Series) -> str:
    filled = _fill_text(s)
    if filled.empty:
        return "не указано"
    return str(filled.value_counts().idxmax())


def _top_share_pct(s: pd.Series) -> float:
    filled = _fill_text(s)
    if filled.empty:
        return 0.0
    vc = filled.value_counts(normalize=True)
    if vc.empty:
        return 0.0
    return float(vc.iloc[0] * 100.0)


def _prepare_backtesting_summary(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    xl = pd.ExcelFile(path)
    if "summary_long" not in xl.sheet_names:
        if "Среднее отклонение" in xl.sheet_names and "Основные группы" in xl.sheet_names:
            avg = pd.read_excel(path, sheet_name="Среднее отклонение")
            main = pd.read_excel(path, sheet_name="Основные группы")
            return avg, main
        if "Сводка по прогнозу" in xl.sheet_names:
            raw = pd.read_excel(path, sheet_name="Сводка по прогнозу", header=None)
            raw = raw.dropna(how="all").reset_index(drop=True)
            header_rows = raw.index[raw.iloc[:, 0].astype("string").eq("Подразделение")].tolist()
            if len(header_rows) >= 2:
                first_header = header_rows[0]
                second_header = header_rows[1]

                main_cols = raw.iloc[first_header].tolist()
                main = raw.iloc[first_header + 1 : second_header - 1].copy()
                main.columns = main_cols
                main = main.dropna(how="all").reset_index(drop=True)

                avg_cols = raw.iloc[second_header].tolist()
                avg = raw.iloc[second_header + 1 :].copy()
                avg.columns = avg_cols
                avg = avg.dropna(how="all").reset_index(drop=True)
                return avg, main
        raise ValueError(f"Не удалось найти нужные листы в {path}")

    src = pd.read_excel(path, sheet_name="summary_long")
    src = src.copy()

    summary = src[
        [
            "Подразделение",
            "Показатель",
            "Погрешность, %",
            "Смещение, гол./п.п.",
            "Средняя абс. ошибка, гол./п.п.",
            "Средний прогноз",
            "Средний факт",
            "Точек",
            "Первый месяц",
            "Последний месяц",
        ]
    ].copy()
    summary = summary.rename(
        columns={
            "Погрешность, %": "Среднее отклонение, %",
            "Смещение, гол./п.п.": "Среднее смещение, гол./п.п.",
            "Средняя абс. ошибка, гол./п.п.": "Средняя абсолютная ошибка, гол./п.п.",
        }
    )
    summary = summary.sort_values(["Подразделение", "Показатель"], kind="mergesort").reset_index(drop=True)

    main = src[src["Показатель"].isin(MAIN_GROUPS)].copy()
    main_pivot = main.pivot(index="Подразделение", columns="Показатель", values="Погрешность, %").reset_index()
    ordered_cols = ["Подразделение"] + [c for c in MAIN_GROUPS if c in main_pivot.columns]
    main_pivot = main_pivot[ordered_cols].copy()
    return summary, main_pivot


def _prepare_lactation_summary(pairs_path: str | Path, *, allowed_subdivisions: list[str] | None = None) -> pd.DataFrame:
    df = pd.read_excel(pairs_path, sheet_name="all_pairs")
    df = df.copy()
    if allowed_subdivisions:
        allow = {str(x).upper().strip() for x in allowed_subdivisions}
        df = df[_normalize_name_series(df["conception_subdivision"]).isin(allow)].copy()
    df["days"] = pd.to_numeric(df["days_from_conception_to_calving"], errors="coerce")
    df = df[df["days"].notna()].copy()
    df["days"] = df["days"].astype(float)
    df["lactation_group"] = _fill_text(df["lactation_group"], default="не указано")
    out = (
        df.groupby("lactation_group", dropna=False)["days"]
        .agg(["count", "mean", "median", "min", "max"])
        .reset_index()
        .rename(
            columns={
                "lactation_group": "Группа лактации",
                "count": "Количество случаев",
                "mean": "Средний срок до отёла, дней",
                "median": "Медианный срок, дней",
                "min": "Минимум, дней",
                "max": "Максимум, дней",
            }
        )
    )
    order = {"0": 0, "1": 1, "2": 2, "3": 3, "4": 4, "5+": 5}
    out["_order"] = out["Группа лактации"].map(order).fillna(99)
    out = out.sort_values("_order", kind="mergesort").drop(columns="_order").reset_index(drop=True)
    return out


def _prepare_gestation_clusters(pairs_path: str | Path, *, allowed_subdivisions: list[str] | None = None) -> pd.DataFrame:
    df = pd.read_excel(pairs_path, sheet_name="all_pairs")
    df = df.copy()
    if allowed_subdivisions:
        allow = {str(x).upper().strip() for x in allowed_subdivisions}
        df = df[_normalize_name_series(df["conception_subdivision"]).isin(allow)].copy()
    df["days"] = pd.to_numeric(df["days_from_conception_to_calving"], errors="coerce")
    df = df[df["days"].notna()].copy()
    df["days"] = df["days"].astype(int)
    df["calving_date"] = pd.to_datetime(df["calving_date"], errors="coerce")
    df["month_calving"] = df["calving_date"].dt.month.map(MONTH_RU)
    df["lactation_group"] = _fill_text(df["lactation_group"], default="не указано")
    df["conception_subdivision"] = _fill_text(df["conception_subdivision"], default="не указано")

    bins = [0, 260, 265, 270, 275, 280, 285, 290, 999]
    labels = ["<260", "260–264", "265–269", "270–274", "275–279", "280–284", "285–289", "290+"]
    df["Кластер срока до отёла"] = pd.cut(
        df["days"],
        bins=bins,
        labels=labels,
        right=False,
        include_lowest=True,
    ).astype("string")

    rows = []
    total = max(len(df), 1)
    for cluster, g in df.groupby("Кластер срока до отёла", dropna=False):
        if pd.isna(cluster):
            continue
        rows.append(
            {
                "Кластер срока до отёла": str(cluster),
                "Количество случаев": int(len(g)),
                "Доля от всех случаев, %": float(len(g) / total * 100.0),
                "Средний срок до отёла, дней": float(g["days"].mean()),
                "Медианный срок, дней": float(g["days"].median()),
                "Преобладающее подразделение": _top_category(g["conception_subdivision"]),
                "Доля этого подразделения в кластере, %": _top_share_pct(g["conception_subdivision"]),
                "Преобладающий месяц отёла": _top_category(g["month_calving"]),
                "Доля этого месяца в кластере, %": _top_share_pct(g["month_calving"]),
                "Преобладающая группа лактации": _top_category(g["lactation_group"]),
                "Доля этой лактации в кластере, %": _top_share_pct(g["lactation_group"]),
            }
        )

    out = pd.DataFrame(rows)
    cluster_order = {label: i for i, label in enumerate(labels)}
    out["_order"] = out["Кластер срока до отёла"].map(cluster_order).fillna(999)
    out = out.sort_values("_order", kind="mergesort").drop(columns="_order").reset_index(drop=True)
    return out


def _prepare_first_week_summary(pairs_path: str | Path, *, allowed_subdivisions: list[str] | None = None) -> pd.DataFrame:
    df = pd.read_excel(pairs_path, sheet_name="all_pairs")
    df = df.copy()
    if allowed_subdivisions:
        allow = {str(x).upper().strip() for x in allowed_subdivisions}
        df = df[_normalize_name_series(df["conception_subdivision"]).isin(allow)].copy()
    df["calving_date"] = pd.to_datetime(df["calving_date"], errors="coerce")
    df = df[df["calving_date"].notna()].copy()
    df["Подразделение"] = _fill_text(df["conception_subdivision"], default="не указано")
    df["Месяц"] = df["calving_date"].dt.to_period("M").astype(str)
    df["Первая неделя"] = df["calving_date"].dt.day <= 7

    month_stats = (
        df.groupby(["Подразделение", "Месяц"], dropna=False)
        .agg(
            Отёлов_за_месяц=("calving_date", "size"),
            Отёлов_в_первые_7_дней=("Первая неделя", "sum"),
        )
        .reset_index()
    )
    month_stats["Доля_в_первую_неделю"] = month_stats["Отёлов_в_первые_7_дней"] / month_stats["Отёлов_за_месяц"]

    summary = (
        month_stats.groupby("Подразделение", dropna=False)
        .agg(
            Месяцев_в_расчёте=("Месяц", "nunique"),
            Среднее_число_отёлов_за_месяц=("Отёлов_за_месяц", "mean"),
            Среднее_число_отёлов_в_первые_7_дней=("Отёлов_в_первые_7_дней", "mean"),
            Средняя_доля_отёлов_в_первую_неделю_проц=("Доля_в_первую_неделю", lambda s: s.mean() * 100.0),
        )
        .reset_index()
        .rename(
            columns={
                "Месяцев_в_расчёте": "Месяцев в расчёте",
                "Среднее_число_отёлов_за_месяц": "Среднее число отёлов за месяц",
                "Среднее_число_отёлов_в_первые_7_дней": "Среднее число отёлов в первые 7 дней месяца",
                "Средняя_доля_отёлов_в_первую_неделю_проц": "Средняя доля отёлов в первые 7 дней месяца, %",
            }
        )
    )
    return summary.sort_values("Подразделение", kind="mergesort").reset_index(drop=True)


def _write_excel(
    *,
    out_path: Path,
    avg_deviation: pd.DataFrame,
    main_groups: pd.DataFrame,
    lactation_summary: pd.DataFrame,
    clusters: pd.DataFrame,
    first_week: pd.DataFrame,
) -> dict[str, int]:
    layout: dict[str, int] = {}
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame({"Сводка по прогнозу": ["Основные группы: среднее отклонение, %"]}).to_excel(
            writer,
            sheet_name="Сводка по прогнозу",
            index=False,
            header=False,
            startrow=0,
        )
        main_groups.to_excel(writer, sheet_name="Сводка по прогнозу", index=False, startrow=1)
        layout["main_groups_header_row"] = 2
        layout["main_groups_first_data_row"] = 3
        layout["main_groups_last_data_row"] = 2 + len(main_groups)

        avg_title_row = len(main_groups) + 4
        pd.DataFrame({"Сводка по прогнозу": ["Все группы: среднее отклонение и среднее смещение"]}).to_excel(
            writer,
            sheet_name="Сводка по прогнозу",
            index=False,
            header=False,
            startrow=avg_title_row - 1,
        )
        avg_deviation.to_excel(writer, sheet_name="Сводка по прогнозу", index=False, startrow=avg_title_row)

        pd.DataFrame({"Срок до отёла": ["Срок до отёла по лактации"]}).to_excel(
            writer,
            sheet_name="Срок до отёла",
            index=False,
            header=False,
            startrow=0,
        )
        lactation_summary.to_excel(writer, sheet_name="Срок до отёла", index=False, startrow=1)
        layout["lact_header_row"] = 2
        layout["lact_first_data_row"] = 3
        layout["lact_last_data_row"] = 2 + len(lactation_summary)

        cluster_title_row = len(lactation_summary) + 4
        pd.DataFrame({"Срок до отёла": ["Кластеры срока до отёла"]}).to_excel(
            writer,
            sheet_name="Срок до отёла",
            index=False,
            header=False,
            startrow=cluster_title_row - 1,
        )
        clusters.to_excel(writer, sheet_name="Срок до отёла", index=False, startrow=cluster_title_row)

        first_week_title_row = cluster_title_row + len(clusters) + 3
        pd.DataFrame({"Срок до отёла": ["Отёлы в первые 7 дней месяца"]}).to_excel(
            writer,
            sheet_name="Срок до отёла",
            index=False,
            header=False,
            startrow=first_week_title_row - 1,
        )
        first_week.to_excel(writer, sheet_name="Срок до отёла", index=False, startrow=first_week_title_row)
        layout["first_week_header_row"] = first_week_title_row + 1
        layout["first_week_first_data_row"] = first_week_title_row + 2
        layout["first_week_last_data_row"] = first_week_title_row + 1 + len(first_week)

    return layout


def _add_charts(workbook_path: Path, *, layout: dict[str, int]) -> None:
    wb = load_workbook(workbook_path)
    if "Графики" in wb.sheetnames:
        del wb["Графики"]
    ws = wb.create_sheet("Графики")

    summary_ws = wb["Сводка по прогнозу"]
    duration_ws = wb["Срок до отёла"]

    chart1 = BarChart()
    chart1.type = "bar"
    chart1.style = 10
    chart1.title = "Среднее отклонение по основным группам"
    chart1.x_axis.title = "Среднее отклонение, %"
    chart1.y_axis.title = "Подразделение"
    max_col = min(summary_ws.max_column, 5)
    data = Reference(
        summary_ws,
        min_col=2,
        max_col=max_col,
        min_row=layout["main_groups_header_row"],
        max_row=layout["main_groups_last_data_row"],
    )
    cats = Reference(
        summary_ws,
        min_col=1,
        min_row=layout["main_groups_first_data_row"],
        max_row=layout["main_groups_last_data_row"],
    )
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    ws.add_chart(chart1, "A1")

    chart2 = BarChart()
    chart2.type = "col"
    chart2.style = 11
    chart2.title = "Средний срок до отёла по лактации"
    chart2.y_axis.title = "Дней"
    chart2.x_axis.title = "Группа лактации"
    data = Reference(
        duration_ws,
        min_col=3,
        min_row=layout["lact_header_row"],
        max_row=layout["lact_last_data_row"],
    )
    cats = Reference(
        duration_ws,
        min_col=1,
        min_row=layout["lact_first_data_row"],
        max_row=layout["lact_last_data_row"],
    )
    chart2.add_data(data, titles_from_data=True)
    chart2.set_categories(cats)
    ws.add_chart(chart2, "J1")

    chart3 = BarChart()
    chart3.type = "col"
    chart3.style = 12
    chart3.title = "Доля отёлов в первые 7 дней месяца"
    chart3.y_axis.title = "%"
    chart3.x_axis.title = "Подразделение"
    data = Reference(
        duration_ws,
        min_col=5,
        min_row=layout["first_week_header_row"],
        max_row=layout["first_week_last_data_row"],
    )
    cats = Reference(
        duration_ws,
        min_col=1,
        min_row=layout["first_week_first_data_row"],
        max_row=layout["first_week_last_data_row"],
    )
    chart3.add_data(data, titles_from_data=True)
    chart3.set_categories(cats)
    ws.add_chart(chart3, "A20")

    exports_dir = Path(workbook_path).parent
    extra_images = [
        (exports_dir / "доказательство_слабого_сигнала_1_сила_признаков.png", "J20", 720),
        (exports_dir / "доказательство_слабого_сигнала_2_точечные_графики.png", "A38", 980),
    ]
    for img_path, anchor, target_width in extra_images:
        if not img_path.exists():
            continue
        try:
            img = XLImage(str(img_path))
            if getattr(img, "width", None):
                ratio = target_width / float(img.width)
                img.width = int(img.width * ratio)
                img.height = int(img.height * ratio)
            ws.add_image(img, anchor)
        except Exception:
            continue

    wb.save(workbook_path)


def build_report(
    *,
    backtesting_report_path: str | Path,
    pairs_path: str | Path,
    output_path: str | Path,
) -> Path:
    allowed_subdivisions = _extract_allowed_subdivisions(backtesting_report_path)
    avg_deviation, main_groups = _prepare_backtesting_summary(backtesting_report_path)
    lactation_summary = _prepare_lactation_summary(pairs_path, allowed_subdivisions=allowed_subdivisions)
    clusters = _prepare_gestation_clusters(pairs_path, allowed_subdivisions=allowed_subdivisions)
    first_week = _prepare_first_week_summary(pairs_path, allowed_subdivisions=allowed_subdivisions)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    layout = _write_excel(
        out_path=output_path,
        avg_deviation=avg_deviation,
        main_groups=main_groups,
        lactation_summary=lactation_summary,
        clusters=clusters,
        first_week=first_week,
    )
    _add_charts(output_path, layout=layout)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--backtesting-report",
        default="exports/subdivision_backtesting_report.xlsx",
        help="Path to raw backtesting report",
    )
    parser.add_argument(
        "--pairs",
        default="exports/conception_to_calving_pairs_enalb.xlsx",
        help="Path to conception-to-calving pairs file",
    )
    parser.add_argument(
        "--output",
        default="exports/отчет_для_руководства.xlsx",
        help="Path to final management report",
    )
    args = parser.parse_args()

    out = build_report(
        backtesting_report_path=args.backtesting_report,
        pairs_path=args.pairs,
        output_path=args.output,
    )
    print(f"saved_excel: {out}")


if __name__ == "__main__":
    main()
