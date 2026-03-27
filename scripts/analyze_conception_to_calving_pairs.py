from __future__ import annotations

import argparse
from pathlib import Path

try:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except Exception:  # pragma: no cover - optional dependency
    matplotlib = None
    plt = None
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference


BIN_EDGES = [-np.inf, 260, 270, 280, 290, 300, 310, 320, 330, 340, np.inf]
BIN_LABELS = [
    "<260",
    "260-269",
    "270-279",
    "280-289",
    "290-299",
    "300-309",
    "310-319",
    "320-329",
    "330-339",
    "340+",
]


def _season_from_month(month: int | float | None) -> str:
    if pd.isna(month):
        return "unknown"
    month_i = int(month)
    if month_i in (12, 1, 2):
        return "winter"
    if month_i in (3, 4, 5):
        return "spring"
    if month_i in (6, 7, 8):
        return "summer"
    return "autumn"


def _lact_group(v: object) -> str:
    n = pd.to_numeric(v, errors="coerce")
    if pd.isna(n):
        return "unknown"
    if n <= 0:
        return "0"
    if n <= 4:
        return str(int(n))
    return "5+"


def _fill_text(s: pd.Series, default: str = "unknown") -> pd.Series:
    return s.astype("string").fillna(default).replace("", default)


def _prepare_frame(path: str | Path) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name="all_pairs")
    df = df.copy()
    df["days_from_conception_to_calving"] = pd.to_numeric(df["days_from_conception_to_calving"], errors="coerce")
    df = df[df["days_from_conception_to_calving"].notna()].copy()
    df["days_from_conception_to_calving"] = df["days_from_conception_to_calving"].astype(int)

    for col in ("conception_date", "pregnancy_confirm_date", "calving_date"):
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce")

    df["gestation_bin_10d"] = pd.cut(
        df["days_from_conception_to_calving"],
        bins=BIN_EDGES,
        labels=BIN_LABELS,
        right=False,
        include_lowest=True,
    ).astype("string")
    df["gestation_tail"] = np.where(
        df["days_from_conception_to_calving"] < 270,
        "short",
        np.where(df["days_from_conception_to_calving"] > 320, "long", "core"),
    )
    df["lactation_group"] = df["lactation_at_conception"].map(_lact_group)
    df["conception_month"] = df["conception_date"].dt.month
    df["calving_month"] = df["calving_date"].dt.month
    df["conception_season"] = df["conception_month"].map(_season_from_month)
    df["calving_season"] = df["calving_month"].map(_season_from_month)
    df["bull_breed_filled"] = _fill_text(df.get("bull_breed", pd.Series(index=df.index, dtype="string")))
    df["semen_type_filled"] = _fill_text(df.get("semen_type", pd.Series(index=df.index, dtype="string")), default="non-sex/unknown")
    df["concept_source_filled"] = _fill_text(df.get("concept_source", pd.Series(index=df.index, dtype="string")))
    df["conception_subdivision_filled"] = _fill_text(df.get("conception_subdivision", pd.Series(index=df.index, dtype="string")))
    df["calving_subdivision_filled"] = _fill_text(df.get("calving_subdivision", pd.Series(index=df.index, dtype="string")))
    df["bull_short_name_filled"] = _fill_text(df.get("bull_short_name", pd.Series(index=df.index, dtype="string")))
    df["technician_filled"] = _fill_text(df.get("technician", pd.Series(index=df.index, dtype="string")))
    return df


def _global_stats(df: pd.DataFrame) -> pd.DataFrame:
    s = df["days_from_conception_to_calving"]
    out = {
        "rows": int(len(df)),
        "mean_days": float(s.mean()),
        "median_days": float(s.median()),
        "std_days": float(s.std()),
        "min_days": int(s.min()),
        "p05_days": float(s.quantile(0.05)),
        "p10_days": float(s.quantile(0.10)),
        "p25_days": float(s.quantile(0.25)),
        "p75_days": float(s.quantile(0.75)),
        "p90_days": float(s.quantile(0.90)),
        "p95_days": float(s.quantile(0.95)),
        "max_days": int(s.max()),
        "short_share_pct": float((s < 270).mean() * 100.0),
        "core_share_pct": float(((s >= 270) & (s <= 320)).mean() * 100.0),
        "long_share_pct": float((s > 320).mean() * 100.0),
    }
    return pd.DataFrame([out])


def _bin_stats(df: pd.DataFrame) -> pd.DataFrame:
    s = df["days_from_conception_to_calving"]
    grouped = (
        df.groupby("gestation_bin_10d", dropna=False)["days_from_conception_to_calving"]
        .agg(["count", "mean", "median", "min", "max", "std"])
        .reset_index()
        .rename(
            columns={
                "gestation_bin_10d": "bin",
                "count": "rows",
                "mean": "mean_days",
                "median": "median_days",
                "min": "min_days",
                "max": "max_days",
                "std": "std_days",
            }
        )
    )
    grouped["share_pct"] = grouped["rows"] / max(len(df), 1) * 100.0
    grouped["cum_share_pct"] = grouped["share_pct"].cumsum()
    return grouped


def _feature_stats(df: pd.DataFrame, feature: str, *, min_rows: int = 1, top_n: int | None = None) -> pd.DataFrame:
    work = df[[feature, "days_from_conception_to_calving"]].copy()
    work[feature] = _fill_text(work[feature])
    work["is_short"] = work["days_from_conception_to_calving"] < 270
    work["is_core"] = work["days_from_conception_to_calving"].between(270, 320, inclusive="both")
    work["is_long"] = work["days_from_conception_to_calving"] > 320

    grouped = work.groupby(feature, dropna=False)["days_from_conception_to_calving"]
    out = grouped.agg(["count", "mean", "median", "std", "min", "max"]).reset_index()
    out = out.rename(
        columns={
            feature: "category",
            "count": "rows",
            "mean": "mean_days",
            "median": "median_days",
            "std": "std_days",
            "min": "min_days",
            "max": "max_days",
        }
    )

    extras = (
        work.groupby(feature, dropna=False)
        .agg(
            short_share_pct=("is_short", lambda s: s.mean() * 100.0),
            core_share_pct=("is_core", lambda s: s.mean() * 100.0),
            long_share_pct=("is_long", lambda s: s.mean() * 100.0),
            p10_days=("days_from_conception_to_calving", lambda s: s.quantile(0.10)),
            p90_days=("days_from_conception_to_calving", lambda s: s.quantile(0.90)),
        )
        .reset_index()
    )
    extras = extras.rename(columns={feature: "category"})
    out = out.merge(extras, on="category", how="left")

    global_mean = float(df["days_from_conception_to_calving"].mean())
    global_short = float((df["days_from_conception_to_calving"] < 270).mean() * 100.0)
    global_long = float((df["days_from_conception_to_calving"] > 320).mean() * 100.0)
    out["delta_mean_days_vs_global"] = out["mean_days"] - global_mean
    out["delta_short_pp_vs_global"] = out["short_share_pct"] - global_short
    out["delta_long_pp_vs_global"] = out["long_share_pct"] - global_long
    out["share_pct"] = out["rows"] / max(len(df), 1) * 100.0
    out = out[out["rows"] >= min_rows].copy()
    out = out.sort_values(["rows", "mean_days"], ascending=[False, True], kind="mergesort")
    if top_n is not None:
        out = out.head(top_n).copy()
    return out.reset_index(drop=True)


def _feature_vs_bins(df: pd.DataFrame, feature: str, *, min_rows: int = 30, top_n: int | None = None) -> pd.DataFrame:
    work = df[[feature, "gestation_bin_10d"]].copy()
    work[feature] = _fill_text(work[feature])
    counts = (
        work.groupby([feature, "gestation_bin_10d"], dropna=False)
        .size()
        .rename("rows")
        .reset_index()
        .rename(columns={feature: "category", "gestation_bin_10d": "bin"})
    )
    totals = counts.groupby("category", dropna=False)["rows"].sum().rename("category_rows").reset_index()
    counts = counts.merge(totals, on="category", how="left")
    counts = counts[counts["category_rows"] >= min_rows].copy()
    counts["share_within_category_pct"] = counts["rows"] / counts["category_rows"] * 100.0
    if top_n is not None:
        keep = totals.sort_values("category_rows", ascending=False).head(top_n)["category"].tolist()
        counts = counts[counts["category"].isin(keep)].copy()
    return counts.sort_values(["category_rows", "category", "bin"], ascending=[False, True, True], kind="mergesort")


def _plot_distribution(df: pd.DataFrame, out_path: Path) -> None:
    if plt is None:
        return
    counts = _bin_stats(df)
    fig, ax = plt.subplots(figsize=(12, 6))
    ax.bar(counts["bin"].astype(str), counts["rows"], color="#4878CF")
    ax.set_title("Распределение срока от зачатия до отела")
    ax.set_xlabel("Интервал, дней")
    ax.set_ylabel("Количество пар")
    ax.tick_params(axis="x", rotation=45)
    fig.tight_layout()
    fig.savefig(out_path, dpi=160)
    plt.close(fig)


def _plot_feature_means(stats: pd.DataFrame, title: str, out_path: Path) -> None:
    if plt is None:
        return
    fig, ax = plt.subplots(figsize=(12, 6))
    ax.bar(stats["category"].astype(str), stats["mean_days"], color="#6ACC64")
    ax.axhline(stats["mean_days"].mean(), color="#333333", linestyle="--", linewidth=1)
    ax.set_title(title)
    ax.set_xlabel("")
    ax.set_ylabel("Средний срок, дней")
    ax.tick_params(axis="x", rotation=45)
    fig.tight_layout()
    fig.savefig(out_path, dpi=160)
    plt.close(fig)


def _plot_tail_mix(stats: pd.DataFrame, title: str, out_path: Path) -> None:
    if plt is None:
        return
    fig, ax = plt.subplots(figsize=(12, 6))
    idx = np.arange(len(stats))
    ax.bar(idx, stats["short_share_pct"], label="short <270", color="#D65F5F")
    ax.bar(idx, stats["core_share_pct"], bottom=stats["short_share_pct"], label="core 270-320", color="#4C72B0")
    ax.bar(
        idx,
        stats["long_share_pct"],
        bottom=stats["short_share_pct"] + stats["core_share_pct"],
        label="long >320",
        color="#DD8452",
    )
    ax.set_xticks(idx)
    ax.set_xticklabels(stats["category"].astype(str), rotation=45, ha="right")
    ax.set_ylabel("Доля внутри категории, %")
    ax.set_title(title)
    ax.legend()
    fig.tight_layout()
    fig.savefig(out_path, dpi=160)
    plt.close(fig)


def _write_excel(
    *,
    df: pd.DataFrame,
    out_path: Path,
    global_summary: pd.DataFrame,
    bin_summary: pd.DataFrame,
    by_lactation: pd.DataFrame,
    by_conception_subdivision: pd.DataFrame,
    by_calving_subdivision: pd.DataFrame,
    by_breed: pd.DataFrame,
    by_semen: pd.DataFrame,
    by_conception_season: pd.DataFrame,
    by_bull: pd.DataFrame,
    by_technician: pd.DataFrame,
    bins_vs_subdivision: pd.DataFrame,
    bins_vs_lactation: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        global_summary.to_excel(writer, sheet_name="summary", index=False)
        bin_summary.to_excel(writer, sheet_name="gestation_bins", index=False)
        by_lactation.to_excel(writer, sheet_name="by_lactation", index=False)
        by_conception_subdivision.to_excel(writer, sheet_name="by_concept_subdiv", index=False)
        by_calving_subdivision.to_excel(writer, sheet_name="by_calving_subdiv", index=False)
        by_breed.to_excel(writer, sheet_name="by_breed", index=False)
        by_semen.to_excel(writer, sheet_name="by_semen_type", index=False)
        by_conception_season.to_excel(writer, sheet_name="by_concept_season", index=False)
        by_bull.to_excel(writer, sheet_name="by_bull_top40", index=False)
        by_technician.to_excel(writer, sheet_name="by_technician_top30", index=False)
        bins_vs_subdivision.to_excel(writer, sheet_name="subdiv_vs_bins", index=False)
        bins_vs_lactation.to_excel(writer, sheet_name="lactation_vs_bins", index=False)
        df.to_excel(
            writer,
            sheet_name="pairs_enriched",
            index=False,
        )


def _add_excel_charts(workbook_path: Path) -> None:
    wb = load_workbook(workbook_path)
    if "charts" in wb.sheetnames:
        del wb["charts"]
    ws = wb.create_sheet("charts")

    bins_ws = wb["gestation_bins"]
    lact_ws = wb["by_lactation"]
    subdiv_ws = wb["by_concept_subdiv"]

    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "Распределение срока от зачатия до отела"
    chart1.y_axis.title = "Количество пар"
    chart1.x_axis.title = "Интервал, дней"
    data = Reference(bins_ws, min_col=2, min_row=1, max_row=bins_ws.max_row)
    cats = Reference(bins_ws, min_col=1, min_row=2, max_row=bins_ws.max_row)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    ws.add_chart(chart1, "A1")

    chart2 = BarChart()
    chart2.type = "col"
    chart2.style = 10
    chart2.title = "Средний срок по группе лактации"
    chart2.y_axis.title = "Средний срок, дней"
    data = Reference(lact_ws, min_col=3, min_row=1, max_row=lact_ws.max_row)
    cats = Reference(lact_ws, min_col=1, min_row=2, max_row=lact_ws.max_row)
    chart2.add_data(data, titles_from_data=True)
    chart2.set_categories(cats)
    ws.add_chart(chart2, "A20")

    max_row_subdiv = min(subdiv_ws.max_row, 13)
    chart3 = BarChart()
    chart3.type = "col"
    chart3.style = 10
    chart3.title = "Средний срок по подразделению зачатия"
    chart3.y_axis.title = "Средний срок, дней"
    data = Reference(subdiv_ws, min_col=3, min_row=1, max_row=max_row_subdiv)
    cats = Reference(subdiv_ws, min_col=1, min_row=2, max_row=max_row_subdiv)
    chart3.add_data(data, titles_from_data=True)
    chart3.set_categories(cats)
    ws.add_chart(chart3, "J1")

    chart4 = BarChart()
    chart4.type = "bar"
    chart4.grouping = "stacked"
    chart4.overlap = 100
    chart4.style = 12
    chart4.title = "Короткие / центральные / длинные случаи по подразделению зачатия"
    chart4.x_axis.title = "Доля внутри категории, %"
    data = Reference(subdiv_ws, min_col=8, max_col=10, min_row=1, max_row=max_row_subdiv)
    cats = Reference(subdiv_ws, min_col=1, min_row=2, max_row=max_row_subdiv)
    chart4.add_data(data, titles_from_data=True)
    chart4.set_categories(cats)
    ws.add_chart(chart4, "J20")

    wb.save(workbook_path)


def analyze_file(input_path: str | Path, *, output_path: str | Path, plot_dir: str | Path) -> dict[str, Path]:
    input_path = Path(input_path)
    output_path = Path(output_path)
    plot_dir = Path(plot_dir)
    plot_dir.mkdir(parents=True, exist_ok=True)

    df = _prepare_frame(input_path)
    global_summary = _global_stats(df)
    bin_summary = _bin_stats(df)
    by_lactation = _feature_stats(df, "lactation_group", min_rows=10)
    by_conception_subdivision = _feature_stats(df, "conception_subdivision_filled", min_rows=10)
    by_calving_subdivision = _feature_stats(df, "calving_subdivision_filled", min_rows=10)
    by_breed = _feature_stats(df, "bull_breed_filled", min_rows=10)
    by_semen = _feature_stats(df, "semen_type_filled", min_rows=10)
    by_conception_season = _feature_stats(df, "conception_season", min_rows=10)
    by_bull = _feature_stats(df, "bull_short_name_filled", min_rows=50, top_n=40)
    by_technician = _feature_stats(df, "technician_filled", min_rows=100, top_n=30)
    bins_vs_subdivision = _feature_vs_bins(df, "conception_subdivision_filled", min_rows=50, top_n=20)
    bins_vs_lactation = _feature_vs_bins(df, "lactation_group", min_rows=10)

    _write_excel(
        df=df,
        out_path=output_path,
        global_summary=global_summary,
        bin_summary=bin_summary,
        by_lactation=by_lactation,
        by_conception_subdivision=by_conception_subdivision,
        by_calving_subdivision=by_calving_subdivision,
        by_breed=by_breed,
        by_semen=by_semen,
        by_conception_season=by_conception_season,
        by_bull=by_bull,
        by_technician=by_technician,
        bins_vs_subdivision=bins_vs_subdivision,
        bins_vs_lactation=bins_vs_lactation,
    )
    _add_excel_charts(output_path)

    plot_distribution = plot_dir / "gestation_distribution_bins.png"
    plot_lactation = plot_dir / "mean_days_by_lactation.png"
    plot_subdivision = plot_dir / "mean_days_by_conception_subdivision.png"
    plot_subdivision_tail = plot_dir / "tail_mix_by_conception_subdivision.png"

    outputs = {"excel": output_path}
    if plt is not None:
        _plot_distribution(df, plot_distribution)
        _plot_feature_means(by_lactation, "Средний срок стельности по группе лактации", plot_lactation)
        _plot_feature_means(
            by_conception_subdivision.head(12),
            "Средний срок стельности по подразделению зачатия",
            plot_subdivision,
        )
        _plot_tail_mix(
            by_conception_subdivision.head(12),
            "Короткие / центральные / длинные случаи по подразделению зачатия",
            plot_subdivision_tail,
        )
        outputs.update(
            {
                "plot_distribution": plot_distribution,
                "plot_lactation": plot_lactation,
                "plot_subdivision": plot_subdivision,
                "plot_subdivision_tail": plot_subdivision_tail,
            }
        )
    return outputs


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default="exports/conception_to_calving_pairs_enalb.xlsx",
        help="Path to conception-to-calving pairs Excel",
    )
    parser.add_argument(
        "--output",
        default="exports/conception_to_calving_pairs_enalb_analysis.xlsx",
        help="Path to output analysis Excel",
    )
    parser.add_argument(
        "--plot-dir",
        default="exports/conception_to_calving_pairs_enalb_plots",
        help="Directory for png plots",
    )
    args = parser.parse_args()

    outputs = analyze_file(args.input, output_path=args.output, plot_dir=args.plot_dir)
    print(f"saved_excel: {outputs['excel']}")
    for name, path in outputs.items():
        if name == "excel":
            continue
        print(f"{name}: {path}")


if __name__ == "__main__":
    main()
